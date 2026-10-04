"""A small deck with one chart of each kind python-pptx writes."""
import sys
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData, BubbleChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

prs = Presentation()
blank = prs.slide_layouts[5]  # title only

def slide(title):
    s = prs.slides.add_slide(blank)
    s.shapes.title.text = title
    return s

def category(kind, title, series, cats=("Q1", "Q2", "Q3", "Q4")):
    s = slide(title)
    data = CategoryChartData()
    data.categories = cats
    for name, values in series:
        data.add_series(name, values)
    frame = s.shapes.add_chart(kind, Inches(1), Inches(1.5), Inches(8), Inches(5), data)
    frame.chart.has_title = True
    frame.chart.chart_title.text_frame.text = title
    return frame

category(XL_CHART_TYPE.COLUMN_CLUSTERED, "Sales by quarter", [("North", (120, 135, 150, 170)), ("South", (80, 95, 110, 105))])
category(XL_CHART_TYPE.LINE_MARKERS, "Visitors", [("Visitors", (1000, 1200, 900, 1500))])
category(XL_CHART_TYPE.PIE, "Share", [("Share", (40, 30, 20, 10))], cats=("A", "B", "C", "D"))
category(XL_CHART_TYPE.BAR_STACKED_100, "Mix", [("Online", (60, 55, 70, 65)), ("Store", (40, 45, 30, 35))])
category(XL_CHART_TYPE.AREA_STACKED, "Stacked area", [("A", (1, 2, 3, 4)), ("B", (2, 2, 2, 2))])
category(XL_CHART_TYPE.RADAR_MARKERS, "Skills", [("Ann", (3, 4, 5, 2)), ("Bo", (4, 2, 3, 5))], cats=("Speed", "Power", "Skill", "Luck"))
category(XL_CHART_TYPE.DOUGHNUT, "Rings", [("Ring", (5, 3, 2))], cats=("X", "Y", "Z"))

s = slide("Scatter")
xy = XyChartData()
ser = xy.add_series("Points")
for x, y in ((1, 2.5), (2, 3.5), (3, 1.25)):
    ser.add_data_point(x, y)
s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(1), Inches(1.5), Inches(8), Inches(5), xy)

s = slide("Bubbles")
bub = BubbleChartData()
ser = bub.add_series("Cities")
for x, y, z in ((1, 10, 5), (2, 20, 15), (3, 15, 10)):
    ser.add_data_point(x, y, z)
s.shapes.add_chart(XL_CHART_TYPE.BUBBLE, Inches(1), Inches(1.5), Inches(8), Inches(5), bub)

prs.save(sys.argv[1])
print("saved", sys.argv[1])
